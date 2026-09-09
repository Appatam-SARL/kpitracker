#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Génère un document Word structuré à partir des fiches BPMN + diagrammes Mermaid."""

from __future__ import annotations

import base64
import re
import zlib
from pathlib import Path

import requests
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
PROCESS_DIR = ROOT / "processus"
ASSETS = ROOT / "word-assets"
OUT = ROOT / "Dossier-BPMN-Agences-Securite-Privee.docx"

PROCESS_FILES = sorted(PROCESS_DIR.glob("*.md"))


def set_cell_shading(cell, hex_color: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), hex_color)
    shading.set(qn("w:val"), "clear")
    cell._tePr = cell._tc.get_or_add_tcPr()
    cell._tc.get_or_add_tcPr().append(shading)


def add_page_number(doc: Document) -> None:
    section = doc.sections[0]
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Dossier BPMN 2.0 – Agences de sécurité privée  |  Page ")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run2 = p.add_run()
    run2._r.append(fld_begin)
    run2._r.append(instr)
    run2._r.append(fld_end)
    run2.font.size = Pt(9)


def style_doc(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    pf = style.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE

    for level, size, color in [
        ("Heading 1", 18, RGBColor(0x0B, 0x3D, 0x5C)),
        ("Heading 2", 14, RGBColor(0x14, 0x5A, 0x86)),
        ("Heading 3", 12, RGBColor(0x1F, 0x6F, 0xA5)),
    ]:
        hs = doc.styles[level]
        hs.font.name = "Calibri"
        hs.font.size = Pt(size)
        hs.font.bold = True
        hs.font.color.rgb = color
        hs.paragraph_format.space_before = Pt(14)
        hs.paragraph_format.space_after = Pt(8)


def add_horizontal_line(paragraph) -> None:
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "145A86")
    pBdr.append(bottom)
    pPr.append(pBdr)


def clean_md_inline(text: str) -> str:
    text = text.replace("\u00a0", " ")
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    return text.strip()


def parse_sections(md: str) -> dict[str, str]:
    """Découpe le markdown en sections ## numérotées + compléments."""
    sections: dict[str, str] = {}
    parts = re.split(r"\n(?=## )", md)
    for part in parts:
        m = re.match(r"##\s+(.+?)\n(.*)", part, re.S)
        if not m:
            continue
        title = m.group(1).strip()
        body = m.group(2).strip()
        sections[title] = body
    return sections


def extract_code_block(body: str, lang: str | None = None) -> str | None:
    if lang:
        pattern = rf"```{lang}\s*\n(.*?)```"
    else:
        pattern = r"```(?:\w*)\s*\n(.*?)```"
    m = re.search(pattern, body, re.S)
    return m.group(1).strip() if m else None


def extract_tables(body: str) -> list[list[list[str]]]:
    tables = []
    lines = body.splitlines()
    i = 0
    while i < len(lines):
        if "|" in lines[i] and i + 1 < len(lines) and re.match(r"^\s*\|?\s*-+", lines[i + 1]):
            rows = []
            while i < len(lines) and "|" in lines[i]:
                if re.match(r"^\s*\|?\s*-+", lines[i]):
                    i += 1
                    continue
                cells = [clean_md_inline(c) for c in lines[i].strip().strip("|").split("|")]
                rows.append(cells)
                i += 1
            if rows:
                tables.append(rows)
            continue
        i += 1
    return tables


def body_without_code_and_tables(body: str) -> str:
    text = re.sub(r"```.*?```", "", body, flags=re.S)
    lines = []
    for line in text.splitlines():
        if "|" in line and re.match(r"^\s*\|", line):
            continue
        if re.match(r"^\s*\|?\s*-{2,}", line):
            continue
        lines.append(line)
    return "\n".join(lines).strip()


def add_paragraphs_from_text(doc: Document, text: str) -> None:
    if not text:
        return
    blocks = re.split(r"\n\s*\n", text)
    for block in blocks:
        block = block.strip()
        if not block:
            continue
        # listes
        if all(re.match(r"^[-*•]\s+|^\d+\.\s+", ln.strip()) or not ln.strip() for ln in block.splitlines() if ln.strip()):
            for ln in block.splitlines():
                ln = ln.strip()
                if not ln:
                    continue
                ln = re.sub(r"^[-*•]\s+", "", ln)
                ln = re.sub(r"^\d+\.\s+", "", ln)
                p = doc.add_paragraph(clean_md_inline(ln), style="List Bullet")
                for run in p.runs:
                    run.font.size = Pt(10)
            continue
        p = doc.add_paragraph(clean_md_inline(block.replace("\n", " ")))
        for run in p.runs:
            run.font.size = Pt(10)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            val = row[c_idx] if c_idx < len(row) else ""
            cell.text = val
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8)
                    if r_idx == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            if r_idx == 0:
                set_cell_shading(cell, "145A86")
            elif r_idx % 2 == 0:
                set_cell_shading(cell, "F2F7FB")
    doc.add_paragraph()


def sanitize_mermaid(src: str) -> str:
    """Normalise le Mermaid pour les moteurs de rendu (Kroki / mermaid.ink)."""
    s = src.replace("<br/>", " ").replace("<br>", " ").replace("<br />", " ")
    # Hexagones {{ }} peu supportés → rectangles
    s = re.sub(r"\{\{([^}]+)\}\}", r"[\1]", s)
    # Parallélogrammes [/ /] → rectangles
    s = re.sub(r"\[/([^/]+)/\]", r"[\1]", s)
    # Guillemets typographiques
    s = s.replace("’", "'").replace("‘", "'").replace("“", '"').replace("”", '"')
    s = s.replace("–", "-").replace("—", "-").replace("…", "...")
    # Labels avec ? ou / dans nœuds non quotés : laisser tel quel si déjà OK
    return s


def render_mermaid_png(mermaid_src: str, out_path: Path, force: bool = False) -> bool:
    """Rendu Mermaid via Kroki puis mermaid.ink en secours."""
    if not force and out_path.exists() and out_path.stat().st_size > 500:
        return True

    cleaned = sanitize_mermaid(mermaid_src)
    mmd_file = out_path.with_suffix(".mmd")
    mmd_file.write_text(cleaned, encoding="utf-8")

    # 1) Kroki POST
    try:
        resp = requests.post(
            "https://kroki.io/mermaid/png",
            data=cleaned.encode("utf-8"),
            headers={"Content-Type": "text/plain"},
            timeout=90,
        )
        if resp.status_code == 200 and len(resp.content) > 500:
            out_path.write_bytes(resp.content)
            return True
        # 2) Kroki GET deflate
        compressed = zlib.compress(cleaned.encode("utf-8"), 9)
        encoded = base64.urlsafe_b64encode(compressed).decode("ascii")
        resp = requests.get(f"https://kroki.io/mermaid/png/{encoded}", timeout=90)
        if resp.status_code == 200 and len(resp.content) > 500:
            out_path.write_bytes(resp.content)
            return True
    except Exception as exc:
        print(f"  [!] Kroki: {exc}")

    # 3) mermaid.ink (base64 brut)
    try:
        b64 = base64.urlsafe_b64encode(cleaned.encode("utf-8")).decode("ascii")
        resp = requests.get(f"https://mermaid.ink/img/{b64}", timeout=90)
        if resp.status_code == 200 and len(resp.content) > 500:
            out_path.write_bytes(resp.content)
            return True
        print(f"  [!] Rendu Mermaid echoue pour {out_path.name} (HTTP {resp.status_code})")
        return False
    except Exception as exc:
        print(f"  [!] Erreur Mermaid {out_path.name}: {exc}")
        return False


def add_cover(doc: Document) -> None:
    for _ in range(3):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("DOSSIER DE CONCEPTION FONCTIONNELLE")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(0x0B, 0x3D, 0x5C)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Modélisation BPMN 2.0")
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(0x14, 0x5A, 0x86)

    prod = doc.add_paragraph()
    prod.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = prod.add_run("Plateforme SaaS – Gestion des agences de sécurité privée\nMarché africain")
    r.font.size = Pt(13)

    doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run(
        "Norme : BPMN 2.0 (OMG)\n"
        "Outils cibles : Camunda Modeler · Bizagi Modeler · Draw.io\n"
        "17 processus métier · Diagrammes Mermaid intégrés\n"
        "Version 1.0"
    )
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    add_horizontal_line(doc.add_paragraph())
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = note.add_run(
        "Document destiné aux équipes Product, Business Analysts,\n"
        "Architectes logiciels et Développeurs SaaS."
    )
    r.italic = True
    r.font.size = Pt(10)
    doc.add_page_break()


def add_toc_placeholder(doc: Document) -> None:
    doc.add_heading("Table des matières", level=1)
    p = doc.add_paragraph(
        "Utilisez dans Word : Références → Table des matières → Table automatique "
        "(les titres Heading 1/2 alimentent le sommaire)."
    )
    for run in p.runs:
        run.italic = True
        run.font.size = Pt(10)

    items = [
        "1. Conventions BPMN 2.0",
        "2. Chaîne de valeur et priorisation MVP",
        "3. Catalogue des 17 processus",
        "4–20. Fiches processus détaillées (sections 1 à 15 + Mermaid + compléments SaaS)",
    ]
    for item in items:
        doc.add_paragraph(item, style="List Number")
    doc.add_page_break()


def add_conventions(doc: Document) -> None:
    conv_path = ROOT / "00-CONVENTIONS-BPMN.md"
    doc.add_heading("1. Conventions BPMN 2.0", level=1)
    if conv_path.exists():
        text = conv_path.read_text(encoding="utf-8")
        sections = parse_sections(text)
        # intro before first ##
        intro = text.split("## ", 1)[0]
        intro = re.sub(r"^#.*\n", "", intro).strip()
        add_paragraphs_from_text(doc, clean_md_inline(intro))
        for title, body in sections.items():
            doc.add_heading(clean_md_inline(title), level=2)
            for table in extract_tables(body):
                add_table(doc, table)
            add_paragraphs_from_text(doc, body_without_code_and_tables(body))
    doc.add_page_break()


def add_mvp_overview(doc: Document) -> None:
    doc.add_heading("2. Chaîne de valeur et priorisation MVP", level=1)
    doc.add_heading("Chaîne de valeur", level=2)
    chain = (
        "Prospection → Contractualisation → Création site\n"
        "Recrutement → Affectation → Prise de service → Pointage / Rondes / Incidents\n"
        "Absences → Remplacement\n"
        "Pointages + Contrat → Facturation\n"
        "Pointages + Absences → Paiement agents\n"
        "Tous processus → Reporting → Audit"
    )
    p = doc.add_paragraph()
    run = p.add_run(chain)
    run.font.name = "Consolas"
    run.font.size = Pt(9)

    doc.add_heading("Priorisation produit", level=2)
    add_table(
        doc,
        [
            ["Vague", "Processus", "Focus"],
            [
                "Vague 1 – Terrain",
                "Affectation, Prise de service, Pointage, Rondes, Incident, Remplacement",
                "Valeur client immédiate + anti-fraude",
            ],
            [
                "Vague 2 – Com / RH",
                "Prospection, Contrat, Site, Recrutement, Absences",
                "Pipeline revenus + vivier agents",
            ],
            [
                "Vague 3 – Finance / Gouv.",
                "Facturation, Paie, Équipements, Pilotage, Reporting, Audit",
                "Cash-flow + conformité",
            ],
        ],
    )
    doc.add_page_break()


def add_catalog(doc: Document) -> None:
    doc.add_heading("3. Catalogue des 17 processus", level=1)
    rows = [["N°", "Processus", "Fichier"]]
    for f in PROCESS_FILES:
        num = f.stem.split("-", 1)[0]
        title = f.stem.split("-", 1)[1].replace("-", " ").title()
        # better title from H1
        md = f.read_text(encoding="utf-8")
        h1 = re.search(r"^#\s+(.+)$", md, re.M)
        name = clean_md_inline(h1.group(1)) if h1 else title
        rows.append([num, name, f.name])
    add_table(doc, rows)
    doc.add_page_break()


SECTION_ORDER_HINTS = [
    "Nom du processus",
    "Objectif",
    "Déclencheur",
    "Acteurs",
    "Préconditions",
    "Description",
    "Tableau",
    "Décisions",
    "exceptions",
    "Données",
    "Notifications",
    "KPI",
    "Recommandations",
    "ASCII",
    "Mermaid",
    "Compléments",
]


def ordered_sections(sections: dict[str, str]) -> list[tuple[str, str]]:
    def rank(title: str) -> tuple[int, str]:
        t = title.lower()
        for i, hint in enumerate(SECTION_ORDER_HINTS):
            if hint.lower() in t:
                return (i, title)
        return (100, title)

    return sorted(sections.items(), key=lambda kv: rank(kv[0]))


def add_process_chapter(doc: Document, path: Path, chapter_num: int) -> None:
    md = path.read_text(encoding="utf-8")
    h1 = re.search(r"^#\s+(.+)$", md, re.M)
    title = clean_md_inline(h1.group(1)) if h1 else path.stem

    print(f"  Processus {chapter_num}: {title}")
    doc.add_heading(f"{chapter_num}. {title}", level=1)

    # métadonnées en tête (lignes **xxx**)
    meta_lines = []
    for line in md.splitlines()[1:25]:
        if line.startswith("## "):
            break
        if line.startswith("**") or line.startswith("---"):
            if line.startswith("**"):
                meta_lines.append(clean_md_inline(line))
    if meta_lines:
        p = doc.add_paragraph()
        run = p.add_run(" | ".join(meta_lines))
        run.font.size = Pt(9)
        run.italic = True
        run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    sections = parse_sections(md)
    for sec_title, body in ordered_sections(sections):
        doc.add_heading(clean_md_inline(sec_title), level=2)
        low = sec_title.lower()

        # Mermaid → image
        if "mermaid" in low:
            mermaid = extract_code_block(body, "mermaid")
            if mermaid:
                img_path = ASSETS / f"{path.stem}-mermaid.png"
                ok = render_mermaid_png(mermaid, img_path)
                if ok and img_path.exists():
                    caption = doc.add_paragraph()
                    r = caption.add_run("Diagramme Mermaid (BPMN approximé en flowchart)")
                    r.bold = True
                    r.font.size = Pt(10)
                    try:
                        doc.add_picture(str(img_path), width=Inches(6.3))
                        last = doc.paragraphs[-1]
                        last.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    except Exception as exc:
                        doc.add_paragraph(f"[Image Mermaid non insérable: {exc}]")
                else:
                    doc.add_paragraph(
                        "Le rendu image Mermaid n’a pas pu être généré (réseau/Kroki). "
                        "Source Mermaid ci-dessous :"
                    )
                    p = doc.add_paragraph()
                    run = p.add_run(mermaid)
                    run.font.name = "Consolas"
                    run.font.size = Pt(7)
            continue

        # ASCII diagram
        if "ascii" in low or "diagramme bpmn" in low:
            ascii_art = extract_code_block(body, None)
            if ascii_art:
                caption = doc.add_paragraph()
                r = caption.add_run("Diagramme BPMN (représentation ASCII)")
                r.bold = True
                r.font.size = Pt(10)
                # Word: mono in chunks to avoid huge paragraphs
                for chunk in ascii_art.splitlines():
                    p = doc.add_paragraph()
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.space_before = Pt(0)
                    run = p.add_run(chunk if chunk else " ")
                    run.font.name = "Consolas"
                    run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
                    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
                    run.font.size = Pt(7)
            # also any leftover text
            rest = body_without_code_and_tables(body)
            if rest:
                add_paragraphs_from_text(doc, rest)
            continue

        # Tables
        tables = extract_tables(body)
        for table in tables:
            add_table(doc, table)

        # JSON payloads
        for m in re.finditer(r"```(?:json|jsonc)?\s*\n(.*?)```", body, re.S):
            p = doc.add_paragraph()
            r = p.add_run("Exemple de payload / schéma :")
            r.italic = True
            r.font.size = Pt(9)
            for line in m.group(1).strip().splitlines():
                lp = doc.add_paragraph()
                lp.paragraph_format.space_after = Pt(0)
                run = lp.add_run(line)
                run.font.name = "Consolas"
                run.font.size = Pt(7)

        rest = body_without_code_and_tables(body)
        # remove json already handled roughly
        rest = re.sub(r"```.*?```", "", rest, flags=re.S)
        add_paragraphs_from_text(doc, rest)

    doc.add_page_break()


def main() -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)
    print("Génération du document Word…")
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    style_doc(doc)
    add_page_number(doc)
    add_cover(doc)
    add_toc_placeholder(doc)
    add_conventions(doc)
    add_mvp_overview(doc)
    add_catalog(doc)

    start_chapter = 4
    for idx, path in enumerate(PROCESS_FILES):
        add_process_chapter(doc, path, start_chapter + idx)

    # annexe
    doc.add_heading(f"{start_chapter + len(PROCESS_FILES)}. Annexe – Utilisation des diagrammes Mermaid", level=1)
    doc.add_paragraph(
        "Les diagrammes Mermaid intégrés dans ce document sont des représentations flowchart "
        "compatibles avec la logique BPMN (XOR, User Task, Service Task, Timer, End). "
        "Pour Camunda Modeler / Bizagi / Draw.io : recréez les Pools, Lanes et éléments BPMN "
        "à partir du tableau des étapes et du diagramme ASCII de chaque fiche."
    )
    doc.add_paragraph(
        "Canaux de notification prioritaires (Afrique) : WhatsApp Business → SMS → Push → Email → IVR."
    )

    doc.save(OUT)
    print(f"OK -> {OUT}")
    print(f"Taille : {OUT.stat().st_size / 1024:.1f} Ko")


if __name__ == "__main__":
    main()
